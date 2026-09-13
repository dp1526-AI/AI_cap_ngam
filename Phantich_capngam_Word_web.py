import streamlit as st
from google import genai
from google.genai import types
import time, os, re, json, io, tempfile
from docx import Document
from docx.shared import Pt

st.set_page_config(
    page_title="AI Phân Tích & Thẩm Định Cáp Ngầm EVNHCMC",
    page_icon="⚡",
    layout="wide"
)

KNOWLEDGE_FILE = "knowledge_history.json"

if not os.path.exists(KNOWLEDGE_FILE):
    with open(KNOWLEDGE_FILE, "w", encoding="utf-8") as f:
        json.dump([], f, ensure_ascii=False, indent=4)

# Quản lý trạng thái thẩm định file
if "qt_valid" not in st.session_state:
    st.session_state.qt_valid = False
    st.session_state.qt_name = ""
    st.session_state.qt_msg = ""
if "bb_valid" not in st.session_state:
    st.session_state.bb_valid = False
    st.session_state.bb_name = ""
    st.session_state.bb_msg = ""

def save_to_knowledge_base(record_info):
    try:
        with open(KNOWLEDGE_FILE, "r+", encoding="utf-8") as f:
            history = json.load(f)
            history.append(record_info)
            f.seek(0)
            json.dump(history[-20:], f, ensure_ascii=False, indent=4)
    except Exception as e:
        st.warning(f"⚠️ Không thể lưu lịch sử: {e}")

def fill_report_bytes(data_dict, template_path="Bao_Cao_mau.docx"):
    if not os.path.exists(template_path):
        return None

    doc = Document(template_path)
    clean_data = {k: str(v).replace("*", "").replace("#", "").replace("$", "").strip() 
                  for k, v in data_dict.items()}

    def replace_in_paragraphs(paragraphs):
        for p in paragraphs:
            for key, value in clean_data.items():
                if key in p.text:
                    p.text = p.text.replace(key, value)
                    for run in p.runs:
                        run.font.name = 'Arial'
                        run.font.size = Pt(11)

    replace_in_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                replace_in_paragraphs(cell.paragraphs)

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

def upload_bytes_to_gemini(client, uploaded_file):
    # Đặt tên file thuần ASCII 7-bit an toàn tuyệt đối cho HTTP Header
    safe_ascii_name = f"doc_{int(time.time() * 1000)}.pdf"
    
    # Đọc trực tiếp byte stream vào io.BytesIO và gán thuộc tính .name bằng tên ASCII
    file_bytes = uploaded_file.getvalue()
    bio = io.BytesIO(file_bytes)
    bio.name = safe_ascii_name

    # Upload trực tiếp từ stream với config display_name
    file = client.files.upload(
        file=bio,
        config=types.UploadFileConfig(
            display_name=safe_ascii_name,
            mime_type="application/pdf"
        )
    )
    
    while file.state.name == "PROCESSING":
        time.sleep(1)
        file = client.files.get(name=file.name)
        
    return file

def check_qt_validity(client, uploaded_file):
    try:
        pdf_file = upload_bytes_to_gemini(client, uploaded_file)
        prompt = """
        Kiểm tra tài liệu PDF này có phải là Quy trình thử nghiệm cáp ngầm 'QT-CT-02' của EVNHCMC hay không.
        Trả về DUY NHẤT định dạng JSON:
        {"is_qt_ct_02": true/false, "reason": "Giải thích ngắn gọn"}
        """
        chat = client.chats.create(
            model="gemini-2.5-flash",
            config=types.GenerateContentConfig(response_mime_type="application/json")
        )
        res = chat.send_message(message=[pdf_file, prompt])
        match = re.search(r'\{.*\}', res.text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        return {"is_qt_ct_02": False, "reason": "Không thể trích xuất JSON"}
    except Exception as e:
        return {"is_qt_ct_02": False, "reason": str(e)}

def check_bb_validity(client, uploaded_file):
    try:
        pdf_file = upload_bytes_to_gemini(client, uploaded_file)
        prompt = """
        Kiểm tra tài liệu PDF này có phải là Biên bản thử nghiệm cáp ngầm (đặc biệt là thí nghiệm CBM, Phóng điện cục bộ PD, hoặc đo đạc cáp ngầm trung/cao thế) hay không.
        Trả về DUY NHẤT định dạng JSON:
        {"is_cable_report": true/false, "reason": "Giải thích ngắn gọn"}
        """
        chat = client.chats.create(
            model="gemini-2.5-flash",
            config=types.GenerateContentConfig(response_mime_type="application/json")
        )
        res = chat.send_message(message=[pdf_file, prompt])
        match = re.search(r'\{.*\}', res.text, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        return {"is_cable_report": False, "reason": "Không thể trích xuất JSON"}
    except Exception as e:
        return {"is_cable_report": False, "reason": str(e)}

# --- GIAO DIỆN CHÍNH ---
st.title("⚡ AI Phân Tích & Thẩm Định Biên Bản Cáp Ngầm - EVNHCMC")
st.caption("Hệ thống thẩm định tức thì tài liệu đầu vào & Đối soát quy trình QT-CT-02")

with st.sidebar:
    st.header("⚙️ Cấu hình hệ thống")
    api_key = st.text_input("Nhập Google Gemini API Key (AQ... hoặc AIzaSy...):", type="password")
    st.divider()
    st.markdown("**Yêu cầu tệp:**")
    st.markdown("- Quy trình: `QT-CT-02`")
    st.markdown("- Biên bản: `Biên bản thử nghiệm / CBM cáp ngầm`")

col1, col2 = st.columns(2)

def clean_api_key(key: str) -> str:
    if not key:
        return ""
    # Chuyển đổi ký tự Cyrillic 'А' (\u0410) sang ký tự Latin 'A' nếu có
    cleaned = key.replace('\u0410', 'A')
    # Loại bỏ khoảng trắng và các ký tự ngoài dải ASCII chuẩn
    cleaned = cleaned.strip().encode('ascii', 'ignore').decode('ascii')
    return cleaned


# --- CỘT 1: UPLOAD VÀ CHECK NGAY FILE QUY TRÌNH ---
with col1:
    st.subheader("1. File Quy trình kỹ thuật")
    uploaded_qt = st.file_uploader("Tải lên Quy trình thử nghiệm (PDF)", type=["pdf"], key="qt_file")
    
    if uploaded_qt:
        if not api_key:
            st.warning("⚠️ Vui lòng nhập API Key ở thanh bên trái trước để hệ thống thẩm định file.")
        else:
            safe_qt_id = f"{uploaded_qt.name}_{uploaded_qt.size}"
            if st.session_state.qt_name != safe_qt_id:
                with st.spinner("🔍 Đang thẩm định file Quy trình QT-CT-02..."):
                    valid_key = clean_api_key(api_key)
                    client_temp = genai.Client(api_key=valid_key)
                    res_qt = check_qt_validity(client_temp, uploaded_qt)
                    if res_qt.get("is_qt_ct_02"):
                        st.session_state.qt_valid = True
                        st.session_state.qt_name = safe_qt_id
                        st.session_state.qt_msg = res_qt.get("reason", "Quy trình hợp lệ.")
                    else:
                        st.session_state.qt_valid = False
                        st.session_state.qt_name = safe_qt_id
                        st.session_state.qt_msg = res_qt.get("reason", "Không phải quy trình QT-CT-02.")
            
            if st.session_state.qt_valid:
                st.success(f"✅ Hợp lệ: {st.session_state.qt_msg}")
            else:
                st.error(f"❌ SAI QUY TRÌNH: {st.session_state.qt_msg}\n\n👉 Vui lòng tải lại đúng file quy trình QT-CT-02!")

# --- CỘT 2: UPLOAD VÀ CHECK NGAY FILE BIÊN BẢN ---
with col2:
    st.subheader("2. File Biên bản hiện trường")
    uploaded_bb = st.file_uploader("Tải lên Biên bản thử nghiệm cáp (PDF)", type=["pdf"], key="bb_file")
    
    if uploaded_bb:
        if not api_key:
            st.warning("⚠️ Vui lòng nhập API Key ở thanh bên trái trước để hệ thống thẩm định file.")
        else:
            safe_bb_id = f"{uploaded_bb.name}_{uploaded_bb.size}"
            if st.session_state.bb_name != safe_bb_id:
                with st.spinner("🔍 Đang thẩm định file Biên bản thử nghiệm cáp ngầm..."):
                    client_temp = genai.Client(api_key=api_key)
                    res_bb = check_bb_validity(client_temp, uploaded_bb)
                    if res_bb.get("is_cable_report"):
                        st.session_state.bb_valid = True
                        st.session_state.bb_name = safe_bb_id
                        st.session_state.bb_msg = res_bb.get("reason", "Biên bản hợp lệ.")
                    else:
                        st.session_state.bb_valid = False
                        st.session_state.bb_name = safe_bb_id
                        st.session_state.bb_msg = res_bb.get("reason", "Không phải biên bản thử nghiệm cáp ngầm.")
            
            if st.session_state.bb_valid:
                st.success(f"✅ Hợp lệ: {st.session_state.bb_msg}")
            else:
                st.error(f"❌ SAI BIÊN BẢN: {st.session_state.bb_msg}\n\n👉 Vui lòng tải lại đúng file Biên bản thử nghiệm cáp ngầm!")

st.write("")

# --- NÚT BẮT ĐẦU CHẠY PHÂN TÍCH ---
can_run = st.session_state.qt_valid and st.session_state.bb_valid

if st.button("🚀 BẮT ĐẦU THẨM ĐỊNH & PHÂN TÍCH", type="primary", use_container_width=True, disabled=not can_run):
    status_box = st.status("Đang tiến hành phân tích và thẩm định đối soát...", expanded=True)
    
    try:
        client = genai.Client(api_key=api_key)

        status_box.write("⏳ Đang chuẩn bị dữ liệu vào bộ xử lý AI Cloud...")
        pdf_qt = upload_bytes_to_gemini(client, uploaded_qt)
        pdf_bb = upload_bytes_to_gemini(client, uploaded_bb)

        system_instruction = """
        Bạn là Chuyên gia Kiểm định Cáp ngầm Cao thế & Trung thế Bậc cao của EVNHCMC.
        Nhiệm vụ: Phân tích chỉ tiêu Phóng điện cục bộ (CBM/PD), đối soát nghiêm ngặt quy trình QT-CT-02
        và thẩm định ĐỘ TRUNG THỰC/CHÍNH XÁC của phần kết luận trên biên bản hiện trường.
        """

        status_box.write("🧠 Đang bóc tách thông số kỹ thuật và thẩm định tính trung thực của kết luận...")
        
        history_context = ""
        if os.path.exists(KNOWLEDGE_FILE):
            with open(KNOWLEDGE_FILE, "r", encoding="utf-8") as f:
                hist = json.load(f)
                if hist:
                    history_context = f"\n[LỊCH SỬ ĐỐI SOÁT ĐỂ ĐẢM BẢO NHẤT QUÁN]:\n{json.dumps(hist[-3:], ensure_ascii=False)}"

        main_prompt = f"""
        Dựa vào Quy trình QT-CT-02 (Tệp 1) và Biên bản Thử nghiệm Phóng điện cục bộ CBM (Tệp 2), hãy thực hiện phân tích chuyên sâu.
        {history_context}

        Hãy xuất thông tin ĐÚNG BẮT BUỘC theo các thẻ sau:

        [THONG_TIN_CHUNG]
        Tên tuyến cáp, chủng loại, chiều dài, điện áp vận hành, đơn vị thực hiện, ngày thử nghiệm.
        [/THONG_TIN_CHUNG]

        [DANH_GIA]
        Phân tích số liệu CBM chi tiết: Điện áp khởi phát PD, Biên độ phóng điện tối đa (pC hoặc dBmV), vị trí điểm PD (m), tần suất xung. So sánh cụ thể với ngưỡng tiêu chuẩn trong QT-CT-02.
        [/DANH_GIA]

        [DANH_GIA_KET_LUAN]
        - Trích dẫn nguyên văn "Nội dung Kết luận" ghi trên Biên bản hiện trường.
        - Đánh giá Chuyên gia: So sánh kết luận gốc đó với dữ liệu kỹ thuật thực tế đo được.
        - Khẳng định rõ ràng: Kết luận trên biên bản là KHỚP (Chính xác) hay KHÔNG KHỚP (Sai lệch/Bỏ sót rủi ro/Chẩn đoán Đạt hay Không Đạt sai quy trình QT-CT-02). Nêu rõ lý do.
        [/DANH_GIA_KET_LUAN]

        [NHAN_DINH]
        Nhận định chuyên môn về mức độ lão hóa cách điện, rủi ro phóng điện hộp nối/đầu cáp, dự báo nguy cơ sự cố.
        [/NHAN_DINH]

        [KHUYEN_NGHI]
        Đề xuất vận hành (Rút ngắn chu kỳ CBM, sửa chữa hộp nối, giảm tải hoặc thay thế).
        [/KHUYEN_NGHI]
        """

        chat = client.chats.create(
            model="gemini-2.5-flash",
            config=types.GenerateContentConfig(system_instruction=system_instruction)
        )
        response = chat.send_message(message=[pdf_qt, pdf_bb, main_prompt])
        text = response.text

        status_box.update(label="Thẩm định hoàn tất thành công!", state="complete", expanded=False)

        def extract_content(tag, full_text):
            pattern = rf"\[{tag}\](.*?)\[/{tag}\]"
            match = re.search(pattern, full_text, re.DOTALL)
            return match.group(1).strip() if match else "Dữ liệu trống"

        data = {
            "{{THONG_TIN_CHUNG}}": extract_content("THONG_TIN_CHUNG", text),
            "{{DANH_GIA}}": extract_content("DANH_GIA", text),
            "{{DANH_GIA_KET_LUAN}}": extract_content("DANH_GIA_KET_LUAN", text),
            "{{NHAN_DINH}}": extract_content("NHAN_DINH", text),
            "{{KHUYEN_NGHI}}": extract_content("KHUYEN_NGHI", text)
        }

        save_to_knowledge_base({
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "file_bb": uploaded_bb.name,
            "evaluation_match": data["{{DANH_GIA_KET_LUAN}}"][:250]
        })

        st.subheader("📑 Kết quả thẩm định & Đối soát kỹ thuật")
        
        with st.expander("📌 1. Thông tin chung tuyến cáp", expanded=True):
            st.write(data["{{THONG_TIN_CHUNG}}"])

        with st.expander("📊 2. Đánh giá số liệu đo CBM so với QT-CT-02", expanded=True):
            st.write(data["{{DANH_GIA}}"])

        with st.expander("🔍 3. Thẩm định tính chuẩn xác của Kết luận Biên bản", expanded=True):
            if "KHÔNG KHỚP" in data["{{DANH_GIA_KET_LUAN}}"].upper():
                st.error(data["{{DANH_GIA_KET_LUAN}}"])
            else:
                st.success(data["{{DANH_GIA_KET_LUAN}}"])

        with st.expander("⚠️ 4. Nhận định rủi ro & Cách điện", expanded=True):
            st.write(data["{{NHAN_DINH}}"])

        with st.expander("💡 5. Khuyến nghị giải pháp vận hành", expanded=True):
            st.write(data["{{KHUYEN_NGHI}}"])

        doc_bytes = fill_report_bytes(data)
        if doc_bytes:
            st.download_button(
                label="📥 TẢI XUỐNG BÁO CÁO HOÀN THIỆN (.DOCX)",
                data=doc_bytes,
                file_name=f"Bao_Cao_CBM_{int(time.time())}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True
            )
        else:
            st.info("ℹ️ Không tìm thấy file `Bao_Cao_mau.docx` trong thư mục chạy để tạo file tải về.")[cite: 4]

    except Exception as e:
        status_box.update(label="Có lỗi phát sinh!", state="error")
        st.error(f"❌ Lỗi hệ thống: {e}")
